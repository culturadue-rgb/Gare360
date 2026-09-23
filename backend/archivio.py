"""
Archivio storico delle gare: Sociale, Cultura, Servizi educativi.

Tre archivi con le STESSE 34 colonne nello stesso ordine. Sono la memoria
dell'azienda: da qui il simulatore ricava la resa tecnica, il livello dei
concorrenti e i ribassi praticati.

DOVE VIVONO I DATI
------------------
La sorgente è intercambiabile, così si può cambiare senza riscrivere l'app:

    FonteExcel   un file .xlsx (oggi)
    FonteDrive   un foglio Google (domani, quando il Drive sarà collegato)

Si sceglie da sola:
    se GARE360_ARCHIVIO_ID è impostata e Drive risponde  -> FonteDrive
    altrimenti                                           -> FonteExcel

Passare dall'una all'altra non tocca il resto del codice: cambiano solo le
variabili d'ambiente. Il file Excel e il foglio Google hanno la stessa
struttura, quindi lo stesso archivio funziona in entrambi i modi.

ATTENZIONE AL PIANO GRATUITO DI RENDER
--------------------------------------
Con FonteExcel il file sta in DATA_DIR, che su Render Free si svuota a ogni
riavvio. Perciò l'app deve sempre permettere di RISCARICARE l'archivio in
Excel (copia di sicurezza) e di ricaricarlo. Finché non c'è il Drive, quella
copia è l'unica cosa che sopravvive con certezza.
"""

from __future__ import annotations

import io
import os
import re
import unicodedata
from pathlib import Path
from typing import Any, Optional

DATA_DIR = Path(os.environ.get("DATA_DIR", Path(__file__).parent / "data"))
FILE_ARCHIVIO = DATA_DIR / "Archivio_Gare360_unificato.xlsx"

# I tre archivi veri. Legenda e Log_correzioni esistono nel file ma l'app non
# li tocca: sono di sola consultazione per le persone.
ARCHIVI = ["Sociale", "Cultura", "Servizi_educativi"]
FOGLI_IGNORATI = {"Legenda", "Log_correzioni"}

# Le 34 colonne, nell'ordine esatto del foglio. L'ordine conta: le scritture
# compongono le righe posizionalmente.
COLONNE = [
    "id_gara", "cig", "stazione_appaltante", "titolo_gara", "settore", "area",
    "Coordinatore_revisione", "scadenza_gara", "stato_gara", "esito_gara",
    "url_cartella", "note", "base_asta", "regione", "comune",
    "punteggio_economico", "punteggio_tecnico", "max_punteggio_tecnico",
    "scarto_tecnico", "ore_lavoro", "data_segnalazione", "data_consegna",
    "relazioni_tecniche", "max_punteggio_economico",
    "punteggio_tecnico_aggiudicatario", "punteggio_economico_aggiudicatario",
    "n_concorrenti", "uscente", "importo_offerto", "nostro_ribasso",
    "modalita_partecipazione", "data_aggiudicazione", "provincia", "origine_dato",
]

# Colonne che il simulatore legge davvero (dalla Legenda del foglio).
COLONNE_SIMULATORE = [
    "esito_gara", "base_asta", "punteggio_economico", "punteggio_tecnico",
    "max_punteggio_tecnico", "scarto_tecnico", "max_punteggio_economico",
    "punteggio_tecnico_aggiudicatario", "punteggio_economico_aggiudicatario",
    "n_concorrenti", "nostro_ribasso",
]

NUMERICHE = {
    "base_asta", "punteggio_economico", "punteggio_tecnico", "max_punteggio_tecnico",
    "scarto_tecnico", "ore_lavoro", "relazioni_tecniche", "max_punteggio_economico",
    "punteggio_tecnico_aggiudicatario", "punteggio_economico_aggiudicatario",
    "n_concorrenti", "importo_offerto", "nostro_ribasso",
}

ESITI = ["Vinta", "Persa", "In attesa", "Non applicabile"]
STATI_GARA = ["Chiusa", "Abbandonata", "In valutazione"]


class ErroreArchivio(Exception):
    """Problema leggendo o scrivendo l'archivio, con messaggio per l'utente."""


# --------------------------------------------------------------------------- #
#  Lettura dei valori                                                          #
# --------------------------------------------------------------------------- #

def _senza_accenti(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def chiave(nome: str) -> str:
    """Normalizza un nome di colonna per confrontarlo (maiuscole, spazi, accenti)."""
    return re.sub(r"\s+", "_", _senza_accenti(str(nome or "")).strip().lower())


def numero(valore: Any) -> Optional[float]:
    """
    Legge un numero scritto all'italiana. Restituisce None se il valore non è
    interpretabile: mai un numero inventato.

        "802.097,23 €"  -> 802097.23
        "16,64%"        -> 16.64
        "2.315,888,89"  -> None   (due virgole: valore rotto, va corretto a mano)
    """
    if valore is None:
        return None
    if isinstance(valore, (int, float)):
        return float(valore)
    testo = str(valore).strip()
    if not testo:
        return None
    testo = testo.replace("€", " ").replace("%", " ").replace("EUR", " ").strip()
    # Si cattura TUTTA la sequenza di cifre e separatori, comprese eventuali
    # virgole in piu': servono a riconoscere i valori rotti. Fermandosi alla
    # prima virgola, "2.315,888,89" sembrerebbe un innocuo 2.315,888.
    m = re.search(r"-?\d[\d.,\s]*", testo)
    if not m:
        return None
    tok = m.group(0).replace(" ", "").rstrip(".,")
    if tok.count(",") > 1:
        return None  # due separatori decimali: il valore e' da correggere a mano
    intero, _, decimale = tok.partition(",")
    gruppi = intero.split(".")
    if len(gruppi) > 1:
        # I punti devono essere separatori di migliaia regolari, altrimenti il
        # valore è stato rovinato (tipico: un importo diventato data o orario).
        if any(not re.fullmatch(r"\d{3}", g) for g in gruppi[1:]):
            return None
    cifre = "".join(gruppi)
    if not re.fullmatch(r"-?\d+", cifre):
        return None
    try:
        return float(cifre + ("." + decimale if decimale else ""))
    except ValueError:
        return None


def percentuale(valore: Any) -> Optional[float]:
    """
    Riporta una percentuale alla convenzione unica dell'app: 16.64 = 16,64%.

    Serve perche' nel foglio convivono due modi di scrivere la stessa cosa:
      - celle formattate "percentuale": il valore memorizzato e' 0,1664
      - celle scritte a mano: "16,64%" oppure 16,64
    Senza uniformare, il simulatore confronterebbe 0,1664 con 16,64 e
    sbaglierebbe di cento volte.

    Regola: un valore minore di 1 e' una frazione (0,0025 -> 0,25%); da 1 in su
    e' gia' una percentuale (1 -> 1%, 16,64 -> 16,64%). Un ribasso espresso come
    frazione pari o superiore a 1 vorrebbe dire "100% di sconto": non esiste.
    """
    v = numero(valore)
    if v is None:
        return None
    return round(v * 100, 4) if -1 < v < 1 and v != 0 else v


def _testo(valore: Any) -> str:
    """
    Porta un valore a testo nella convenzione italiana.

    I numeri letti da Excel arrivano come float e Python li scriverebbe col
    PUNTO decimale ("0.166379", "802097.23"). Riletti da numero(), che segue la
    convenzione italiana, quel punto verrebbe scambiato per separatore di
    migliaia: "52.958" diventerebbe 52958, mille volte tanto, senza alcun
    errore visibile. Per questo qui si scrive sempre con la VIRGOLA.
    """
    if valore is None:
        return ""
    if isinstance(valore, bool):
        return "Si" if valore else "No"
    if isinstance(valore, int):
        return str(valore)
    if isinstance(valore, float):
        if valore.is_integer():
            return str(int(valore))
        # :.10g evita la notazione esponenziale sui valori del nostro dominio
        return f"{valore:.10g}".replace(".", ",")
    return str(valore).strip()


# --------------------------------------------------------------------------- #
#  Calcoli automatici sulle righe                                              #
# --------------------------------------------------------------------------- #

def completa(riga: dict) -> dict:
    """
    Riempie i due campi ricavabili da altri, SENZA sovrascrivere ciò che c'è già.

      scarto_tecnico  = punteggio tecnico del vincitore - il nostro (0 se vinta)
      nostro_ribasso  = 1 - importo_offerto / base_asta, in percentuale

    Se i dati di partenza mancano o sono illeggibili il campo resta vuoto: si
    dichiara la lacuna invece di inventare un numero.
    """
    # Tutti i valori di una riga sono TESTO, anche quelli calcolati qui: il
    # resto del codice (ricerca, scrittura, confronti) conta su questo.
    riga = dict(riga)

    if not _testo(riga.get("scarto_tecnico")):
        esito = _testo(riga.get("esito_gara")).lower()
        nostro = numero(riga.get("punteggio_tecnico"))
        vincitore = numero(riga.get("punteggio_tecnico_aggiudicatario"))
        if esito == "vinta" and nostro is not None:
            riga["scarto_tecnico"] = _testo(0)
        elif nostro is not None and vincitore is not None:
            riga["scarto_tecnico"] = _testo(round(vincitore - nostro, 2))

    if _testo(riga.get("nostro_ribasso")):
        # Gia' presente: lo si riporta comunque alla convenzione unica, perche'
        # nel foglio alcune celle sono frazioni e altre percentuali.
        p = percentuale(riga["nostro_ribasso"])
        if p is not None:
            riga["nostro_ribasso"] = _testo(round(p, 2))
    else:
        offerto = numero(riga.get("importo_offerto"))
        base = numero(riga.get("base_asta"))
        if offerto is not None and base:
            riga["nostro_ribasso"] = _testo(round((1 - offerto / base) * 100, 2))

    return riga


# --------------------------------------------------------------------------- #
#  Sorgenti dei dati                                                           #
# --------------------------------------------------------------------------- #

class Fonte:
    """Interfaccia comune a Excel e Drive."""

    nome = "?"

    def disponibile(self) -> bool: raise NotImplementedError
    def leggi(self, archivio: str) -> list[dict]: raise NotImplementedError
    def scrivi(self, archivio: str, righe: list[dict]) -> None: raise NotImplementedError
    def descrizione(self) -> str: return self.nome


class FonteExcel(Fonte):
    """L'archivio come file .xlsx in DATA_DIR."""

    nome = "excel"

    def __init__(self, percorso: Path = FILE_ARCHIVIO):
        self.percorso = Path(percorso)

    def disponibile(self) -> bool:
        return self.percorso.is_file()

    def descrizione(self) -> str:
        return f"file Excel {self.percorso.name}" if self.disponibile() else "file Excel non ancora caricato"

    def _apri(self):
        try:
            import openpyxl
        except ImportError as e:
            raise ErroreArchivio(
                "Manca la libreria openpyxl, necessaria a leggere i file Excel."
            ) from e
        if not self.disponibile():
            raise ErroreArchivio(
                "L'archivio non è ancora stato caricato. Vai in Impostazioni e "
                "carica il file Archivio_Gare360_unificato.xlsx."
            )
        return openpyxl.load_workbook(self.percorso, data_only=True)

    def leggi(self, archivio: str) -> list[dict]:
        wb = self._apri()
        foglio = trova_foglio(wb.sheetnames, archivio)
        if not foglio:
            raise _errore_foglio_mancante(archivio, wb.sheetnames)
        ws = wb[foglio]
        r_intestazioni = _riga_intestazioni(ws, foglio)
        intestazioni = [_testo(c.value) for c in ws[r_intestazioni]]
        _verifica_intestazioni(archivio, intestazioni)

        righe = []
        for r in range(r_intestazioni + 1, ws.max_row + 1):
            valori = [ws.cell(r, c).value for c in range(1, len(COLONNE) + 1)]
            if not _testo(valori[0]):
                continue  # riga predisposta ma vuota
            righe.append({COLONNE[i]: _testo(v) for i, v in enumerate(valori)})
        wb.close()
        return righe

    def scrivi(self, archivio: str, righe: list[dict]) -> None:
        import openpyxl
        wb = self._apri()
        foglio = trova_foglio(wb.sheetnames, archivio)
        if not foglio:
            foglio = archivio
            wb.create_sheet(foglio)
            wb[foglio].append(COLONNE)
        ws = wb[foglio]
        r_intestazioni = _riga_intestazioni(ws, foglio)
        # Si riscrive solo il corpo: l'intestazione, e l'eventuale titolo che le
        # sta sopra, restano dove sono.
        ws.delete_rows(r_intestazioni + 1, max(ws.max_row - r_intestazioni, 0))
        for riga in righe:
            ws.append([riga.get(c, "") for c in COLONNE])
        self.percorso.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.percorso)
        wb.close()


class FonteDrive(Fonte):
    """
    L'archivio come foglio Google. Non ancora attiva: si accende impostando
    GARE360_ARCHIVIO_ID e le credenziali Google. Il codice che la userà è già
    pronto in drive.py; qui c'è solo l'innesto, così il passaggio sarà una
    variabile d'ambiente e non una riscrittura.
    """

    nome = "drive"

    def __init__(self, foglio_id: str):
        self.foglio_id = foglio_id

    def disponibile(self) -> bool:
        if not self.foglio_id:
            return False
        try:
            import drive
            return drive.configurato()
        except Exception:
            return False

    def descrizione(self) -> str:
        return f"foglio Google {self.foglio_id}"

    def leggi(self, archivio: str) -> list[dict]:
        import drive
        intestazioni, righe = drive.leggi_dizionari(self.foglio_id, archivio)
        _verifica_intestazioni(archivio, intestazioni)
        return [{c: _testo(r.get(c, "")) for c in COLONNE} for r in righe
                if _testo(r.get("id_gara"))]

    def scrivi(self, archivio: str, righe: list[dict]) -> None:
        import drive
        drive.verifica_intestazioni(self.foglio_id, archivio, COLONNE)
        ultima = drive.colonna_lettera(len(COLONNE))
        drive.scrivi_intervallo(
            self.foglio_id, f"{archivio}!A2:{ultima}{len(righe) + 1}",
            [[riga.get(c, "") for c in COLONNE] for riga in righe],
        )


# --------------------------------------------------------------------------- #
#  Riconoscere i fogli quando il nome non e' esattamente quello atteso         #
# --------------------------------------------------------------------------- #

# Parole che nel nome di un foglio non distinguono niente: quasi ogni foglio di
# un archivio gare si chiama "...Gare...". Toglierle lascia solo cio' che conta.
_RUMORE = {"gare", "gara", "archivio", "archivi", "storico", "storici", "demo",
           "definitiva", "definitivo", "foglio", "dati", "elenco", "tabella",
           "copia", "nuovo", "nuova", "aggiornato", "aggiornata"}


def _parole(nome: str) -> set[str]:
    """
    Le parole significative di un nome, senza accenti, punteggiatura e rumore.

    Si spezza su tutto cio' che non e' lettera o cifra: chiave() normalizza gli
    accenti e le maiuscole ma lascia stare punti e parentesi, e senza toglierli
    "Serv." resterebbe una parola diversa da "serv".
    """
    return {p for p in re.split(r"[^a-z0-9]+", chiave(nome)) if p and p not in _RUMORE}


def _stessa_parola(a: str, b: str) -> bool:
    """
    Due parole indicano la stessa cosa? Uguali, oppure una e' l'abbreviazione
    dell'altra: "serv" sta per "servizi", "cultura" per "culturale".

    Si richiede che la piu' corta sia l'INIZIO della piu' lunga, e lunga almeno
    quattro lettere. Non basta che comincino uguale: "socio" e "sociale"
    condividono quattro lettere ma nessuna delle due inizia l'altra, e
    accomunarle vorrebbe dire scambiare il socio-sanitario per il sociale.
    """
    if a == b:
        return True
    corta, lunga = (a, b) if len(a) <= len(b) else (b, a)
    return len(corta) >= 4 and lunga.startswith(corta)


def trova_foglio(nomi: list[str], archivio: str) -> Optional[str]:
    """
    Quale foglio contiene l'archivio richiesto.

    Prima si cerca il nome esatto. Se non c'e', si accetta un foglio che
    contenga tutte le parole dell'archivio: "Sociale_Gare" e "Demo_Gare_Cultura"
    sono evidentemente il Sociale e la Cultura, e rifiutarli obbligherebbe
    soltanto a rinominare un foglio a mano per far contento il programma.

    Se i candidati sono due ci si ferma: indovinare quale dei due e' quello buono
    significherebbe leggere l'archivio sbagliato senza dirlo.
    """
    for n in nomi:
        if chiave(n) == chiave(archivio):
            return n
    cercate = _parole(archivio)
    if not cercate:
        return None
    candidati = [n for n in nomi
                 if all(any(_stessa_parola(c, p) for p in _parole(n)) for c in cercate)]
    if len(candidati) == 1:
        return candidati[0]
    if len(candidati) > 1:
        raise ErroreArchivio(
            f"Nel file ci sono piu' fogli che sembrano l'archivio «{archivio}»: "
            f"{', '.join(candidati)}. Rinominane uno, o lascia solo quello giusto: "
            "scegliere da solo vorrebbe dire leggere forse l'archivio sbagliato."
        )
    return None


def _errore_foglio_mancante(archivio: str, nomi: list[str]) -> ErroreArchivio:
    """Detto in modo che si capisca cosa fare, non solo che e' andata male."""
    riconosciuti = []
    for a in ARCHIVI:
        try:
            trovato = trova_foglio(nomi, a)
        except ErroreArchivio:
            trovato = None
        if trovato:
            riconosciuti.append(f"{trovato} -> {a}")
    coda = ("Fogli riconosciuti: " + "; ".join(riconosciuti) + ". "
            if riconosciuti else "Nessun foglio di questo file e' stato riconosciuto. ")
    return ErroreArchivio(
        f"Nel file Excel non c'e' nessun foglio per l'archivio «{archivio}». "
        f"Fogli presenti: {', '.join(nomi)}. " + coda +
        "Controlla di aver caricato Archivio_Gare360_unificato.xlsx e non un altro file."
    )


RIGHE_INTESTAZIONE = 20     # fin dove si cerca la riga delle intestazioni


def _riga_intestazioni(ws, foglio: str) -> int:
    """
    In quale riga stanno le intestazioni.

    Quasi sempre la prima, ma un foglio puo' avere sopra un titolo, una riga
    vuota o una nota, e pretendere la riga 1 significherebbe rifiutare un file
    giusto per una riga di troppo. Si cerca quindi la prima riga che comincia
    con "id_gara", entro le prime RIGHE_INTESTAZIONE.
    """
    for r in range(1, min(ws.max_row, RIGHE_INTESTAZIONE) + 1):
        if chiave(_testo(ws.cell(r, 1).value)) == "id_gara":
            return r

    # Non trovata: meglio dire cosa c'e' davvero nel foglio, perche' quasi
    # sempre vuol dire che il file caricato non e' l'archivio.
    visto = "Il foglio e' vuoto. "
    for r in range(1, min(ws.max_row, RIGHE_INTESTAZIONE) + 1):
        contenuto = [t for t in (_testo(ws.cell(r, c).value) for c in range(1, 7)) if t]
        if contenuto:
            visto = f"La prima riga con qualcosa dentro e' la {r} e contiene: " + \
                    ", ".join(contenuto) + ". "
            break
    raise ErroreArchivio(
        f"Nel foglio «{foglio}» non trovo la riga delle intestazioni: "
        f"ho cercato una casella «id_gara» nelle prime {RIGHE_INTESTAZIONE} righe "
        f"e non c'e'. " + visto +
        "Quasi sempre questo vuol dire che il file caricato non e' "
        "Archivio_Gare360_unificato.xlsx ma un altro foglio di calcolo."
    )


def _verifica_intestazioni(archivio: str, presenti: list[str]) -> None:
    """
    Le colonne sono ancora quelle attese? Se no ci si ferma dicendo quale non
    torna, invece di scrivere i valori nelle caselle sbagliate.
    """
    attese = [chiave(c) for c in COLONNE]
    trovate = [chiave(c) for c in presenti]
    for i, atteso in enumerate(attese):
        trovato = trovate[i] if i < len(trovate) else ""
        if trovato != atteso:
            raise ErroreArchivio(
                f"Nell'archivio «{archivio}» la colonna {i + 1} dovrebbe chiamarsi "
                f"«{COLONNE[i]}» e invece è «{presenti[i] if i < len(presenti) else '(mancante)'}». "
                "Lettura e scrittura sono bloccate per non rovinare i dati: "
                "rimetti a posto l'intestazione nel file."
            )


def fonte() -> Fonte:
    """La sorgente attiva: Drive se configurato, altrimenti il file Excel."""
    foglio = os.environ.get("GARE360_ARCHIVIO_ID", "").strip()
    if foglio:
        f = FonteDrive(foglio)
        if f.disponibile():
            return f
    return FonteExcel()


# --------------------------------------------------------------------------- #
#  Operazioni sull'archivio                                                    #
# --------------------------------------------------------------------------- #

def elenco(archivio: str, completa_calcoli: bool = True) -> list[dict]:
    if archivio not in ARCHIVI:
        raise ErroreArchivio(f"Archivio sconosciuto: «{archivio}». Validi: {', '.join(ARCHIVI)}.")
    righe = fonte().leggi(archivio)
    return [completa(r) for r in righe] if completa_calcoli else righe


def tutte() -> dict[str, list[dict]]:
    """Tutti e tre gli archivi. Un archivio illeggibile non blocca gli altri."""
    fuori: dict[str, Any] = {}
    for a in ARCHIVI:
        try:
            fuori[a] = elenco(a)
        except ErroreArchivio as e:
            fuori[a] = {"errore": str(e)}
    return fuori


def aggiungi(archivio: str, riga: dict) -> dict:
    """Aggiunge una gara. L'id deve essere nuovo: due gare con lo stesso id
    renderebbero impossibile modificarne una sola."""
    riga = completa({c: _testo(riga.get(c, "")) for c in COLONNE})
    if not riga["id_gara"]:
        raise ErroreArchivio("Serve un id_gara: è il codice con cui la gara viene ritrovata.")
    righe = fonte().leggi(archivio)
    if any(r["id_gara"] == riga["id_gara"] for r in righe):
        raise ErroreArchivio(
            f"Nell'archivio «{archivio}» esiste già una gara con id «{riga['id_gara']}». "
            "Usa un id diverso (per i lotti la convenzione è -L1, -L2…)."
        )
    righe.append(riga)
    fonte().scrivi(archivio, righe)
    return riga


def aggiorna(archivio: str, id_gara: str, campi: dict) -> dict:
    """Modifica una gara. Ogni campo è correggibile a mano, anche dopo il salvataggio."""
    righe = fonte().leggi(archivio)
    for i, r in enumerate(righe):
        if r["id_gara"] == id_gara:
            aggiornata = {**r, **{c: _testo(v) for c, v in campi.items() if c in COLONNE}}
            righe[i] = completa(aggiornata)
            fonte().scrivi(archivio, righe)
            return righe[i]
    raise ErroreArchivio(f"Nell'archivio «{archivio}» non c'è nessuna gara con id «{id_gara}».")


def elimina(archivio: str, id_gara: str) -> None:
    """Cancella una gara. Chi chiama deve aver già chiesto conferma all'utente."""
    righe = fonte().leggi(archivio)
    restanti = [r for r in righe if r["id_gara"] != id_gara]
    if len(restanti) == len(righe):
        raise ErroreArchivio(f"Nell'archivio «{archivio}» non c'è nessuna gara con id «{id_gara}».")
    fonte().scrivi(archivio, restanti)


def cerca(archivio: str, testo: str = "", regione: str = "", esito: str = "",
          anno: str = "") -> list[dict]:
    """Ricerca libera più i tre filtri chiesti: regione, esito, anno."""
    righe = elenco(archivio)
    t = chiave(testo)
    if t:
        righe = [r for r in righe if t in chiave(" ".join(_testo(v) for v in r.values()))]
    if regione:
        righe = [r for r in righe if chiave(r.get("regione")) == chiave(regione)]
    if esito:
        righe = [r for r in righe if chiave(r.get("esito_gara")) == chiave(esito)]
    if anno:
        righe = [r for r in righe if anno in " ".join(
            str(r.get(c, "")) for c in ("scadenza_gara", "data_aggiudicazione", "data_consegna"))]
    return righe


def valori_filtri(archivio: str) -> dict[str, list[str]]:
    """I valori realmente presenti, per riempire i menu a tendina dei filtri."""
    righe = elenco(archivio)
    anni = set()
    for r in righe:
        for c in ("scadenza_gara", "data_aggiudicazione", "data_consegna"):
            for a in re.findall(r"(20\d{2})", str(r.get(c, ""))):
                anni.add(a)
    return {
        "regioni": sorted({r["regione"] for r in righe if r.get("regione")}),
        "esiti": sorted({r["esito_gara"] for r in righe if r.get("esito_gara")}),
        "anni": sorted(anni, reverse=True),
    }


# --------------------------------------------------------------------------- #
#  Copia di sicurezza                                                          #
# --------------------------------------------------------------------------- #

def esporta_excel() -> bytes:
    """
    Tutti e tre gli archivi in un .xlsx con la stessa struttura dell'originale.
    Finché i dati non stanno su Drive questa è l'unica copia che sopravvive a
    un riavvio di Render: va scaricata spesso.
    """
    try:
        import openpyxl
    except ImportError as e:
        raise ErroreArchivio("Manca la libreria openpyxl, necessaria a creare il file Excel.") from e

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for a in ARCHIVI:
        ws = wb.create_sheet(a)
        ws.append(COLONNE)
        try:
            for r in elenco(a):
                ws.append([r.get(c, "") for c in COLONNE])
        except ErroreArchivio:
            pass  # un archivio illeggibile non deve impedire di salvare gli altri
    flusso = io.BytesIO()
    wb.save(flusso)
    wb.close()
    return flusso.getvalue()


def importa_excel(contenuto: bytes) -> dict[str, int]:
    """
    Carica il file Excel dell'archivio, dopo averlo controllato.
    Restituisce quante righe ha trovato per archivio, così l'utente vede subito
    se il file era quello giusto.
    """
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    provvisorio = DATA_DIR / "_archivio_in_verifica.xlsx"
    provvisorio.write_bytes(contenuto)

    try:
        prova = FonteExcel(provvisorio)
        conteggi = {}
        for a in ARCHIVI:
            conteggi[a] = len(prova.leggi(a))   # solleva se le colonne non tornano
    except ErroreArchivio:
        provvisorio.unlink(missing_ok=True)
        raise
    except Exception as e:
        provvisorio.unlink(missing_ok=True)
        raise ErroreArchivio(f"Il file non è un Excel leggibile: {e}") from e

    provvisorio.replace(FILE_ARCHIVIO)
    return conteggi


def stato() -> dict:
    """Riassunto per l'app: da dove legge, e quante gare ci sono."""
    f = fonte()
    info: dict[str, Any] = {"fonte": f.nome, "descrizione": f.descrizione(),
                            "disponibile": f.disponibile()}
    if not f.disponibile():
        return info
    conteggi = {}
    for a in ARCHIVI:
        try:
            conteggi[a] = len(elenco(a))
        except ErroreArchivio as e:
            conteggi[a] = str(e)
    info["gare"] = conteggi
    return info
