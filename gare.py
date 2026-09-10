"""
Gestione delle gare in lavorazione.

Ogni gara è una cartella in DATA_DIR/gare/<id>/ con:
  gara.json        dati, stato, chat, valutazione, dati per il simulatore
  documenti/       i file caricati (originali)
  testi/<doc>.txt  il testo estratto da ogni documento (usato come contesto AI)

La memoria della gara (chat, analisi, dati estratti) resta attiva per
MEMORIA_GIORNI dall'ultima attività; dopo, viene archiviata: rimane
consultabile ma non è più passata all'assistente finché non la riattivi.
"""

from __future__ import annotations

import json
import os
import re
import uuid
from datetime import datetime, timedelta
from pathlib import Path

DATA_DIR = Path(os.environ.get("DATA_DIR", Path(__file__).parent / "data"))
GARE_DIR = DATA_DIR / "gare"

MEMORIA_GIORNI = 15
SETTORI = ["Cultura", "Sociale", "Altro"]
STATI = ["Da valutare", "In analisi", "GO", "NO GO", "In preparazione", "Presentata", "Archiviata"]
STATI_CONCLUSI = {"Presentata", "Archiviata", "NO GO"}
CATEGORIE_DOC = ["bando", "disciplinare", "capitolato", "criteri di valutazione", "allegato tecnico",
                 "chiarimenti", "documento economico", "documento amministrativo", "bozza interna", "altro"]
MAX_CHARS_DOC = 60_000        # testo tenuto per documento
MAX_CHARS_CONTESTO = 160_000  # testo totale passato all'assistente


def _now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


# ---------------------------------------------------------------------------
# Persistenza
# ---------------------------------------------------------------------------

def _dir(gid: str) -> Path:
    return GARE_DIR / gid


def _carica_raw(gid: str) -> dict:
    p = _dir(gid) / "gara.json"
    if not p.exists():
        raise KeyError(gid)
    return json.loads(p.read_text(encoding="utf-8"))


def _salva(g: dict) -> dict:
    d = _dir(g["id"])
    d.mkdir(parents=True, exist_ok=True)
    (d / "gara.json").write_text(json.dumps(g, ensure_ascii=False, indent=2), encoding="utf-8")
    return g


def crea(titolo: str, ente: str = "", settore: str = "Altro", scadenza: str | None = None,
         note: str = "", base_asta: float | None = None) -> dict:
    g = {
        "id": uuid.uuid4().hex[:10],
        "titolo": titolo.strip(),
        "ente": ente.strip(),
        "settore": settore if settore in SETTORI else "Altro",
        "scadenza": scadenza or None,
        "stato": "Da valutare",
        "note": note,
        "base_asta": base_asta,
        "creata": _now(),
        "aggiornata": _now(),
        "ultima_attivita": _now(),
        "memoria_attiva": True,
        "documenti": [],
        "chat": [],
        "valutazione": None,
        "dati_simulatore": None,
        "info_estratte": None,
    }
    return _salva(g)


def carica(gid: str) -> dict:
    g = _carica_raw(gid)
    return _controlla_memoria(g)


def aggiorna(gid: str, campi: dict) -> dict:
    g = _carica_raw(gid)
    for k in ("titolo", "ente", "settore", "scadenza", "stato", "note", "base_asta", "valutazione", "dati_simulatore", "info_estratte"):
        if k in campi and campi[k] is not None or (k in campi and k in ("scadenza", "base_asta")):
            g[k] = campi[k]
    if g["settore"] not in SETTORI:
        g["settore"] = "Altro"
    if g["stato"] not in STATI:
        g["stato"] = "Da valutare"
    g["aggiornata"] = _now()
    return _salva(g)


def elimina(gid: str) -> None:
    import shutil
    d = _dir(gid)
    if d.exists():
        shutil.rmtree(d)


def tocca(g: dict) -> dict:
    """Registra attività: rinnova la finestra di memoria."""
    g["ultima_attivita"] = _now()
    g["aggiornata"] = g["ultima_attivita"]
    g["memoria_attiva"] = True
    return _salva(g)


def _controlla_memoria(g: dict) -> dict:
    """Se sono passati più di MEMORIA_GIORNI dall'ultima attività, archivia la memoria."""
    if g.get("memoria_attiva"):
        ultima = datetime.fromisoformat(g.get("ultima_attivita") or g["creata"])
        if datetime.now() - ultima > timedelta(days=MEMORIA_GIORNI):
            g["memoria_attiva"] = False
            _salva(g)
    return g


def riattiva_memoria(gid: str) -> dict:
    return tocca(_carica_raw(gid))


# ---------------------------------------------------------------------------
# Elenco, scadenze, urgenza
# ---------------------------------------------------------------------------

def giorni_mancanti(scadenza: str | None):
    if not scadenza:
        return None
    try:
        dt = datetime.fromisoformat(scadenza)
    except ValueError:
        return None
    return (dt.date() - datetime.now().date()).days


def urgenza(scadenza: str | None, stato: str = "") -> str:
    """'scaduta' | 'alta' | 'media' | 'bassa' | 'nessuna'"""
    if stato in STATI_CONCLUSI:
        return "nessuna"
    gg = giorni_mancanti(scadenza)
    if gg is None:
        return "nessuna"
    if gg < 0:
        return "scaduta"
    if gg <= 3:
        return "alta"
    if gg <= 7:
        return "media"
    return "bassa"


def riassunto(g: dict) -> dict:
    """Versione leggera per elenchi e dashboard."""
    return {
        "id": g["id"], "titolo": g["titolo"], "ente": g["ente"], "settore": g["settore"],
        "scadenza": g.get("scadenza"), "stato": g["stato"], "aggiornata": g["aggiornata"],
        "ultima_attivita": g.get("ultima_attivita"), "memoria_attiva": g.get("memoria_attiva", True),
        "giorni_mancanti": giorni_mancanti(g.get("scadenza")),
        "urgenza": urgenza(g.get("scadenza"), g["stato"]),
        "n_documenti": len(g.get("documenti", [])),
        "n_messaggi": len(g.get("chat", [])),
        "ha_valutazione": bool(g.get("valutazione")),
        "ha_dati_simulatore": bool(g.get("dati_simulatore")),
        "base_asta": g.get("base_asta"),
    }


def elenco(settore: str | None = None, stato: str | None = None, concluse: bool | None = None) -> list[dict]:
    out = []
    if GARE_DIR.exists():
        for d in GARE_DIR.iterdir():
            if (d / "gara.json").exists():
                try:
                    g = _controlla_memoria(json.loads((d / "gara.json").read_text(encoding="utf-8")))
                except (json.JSONDecodeError, KeyError):
                    continue
                if settore and g["settore"] != settore:
                    continue
                if stato and g["stato"] != stato:
                    continue
                if concluse is True and g["stato"] not in STATI_CONCLUSI:
                    continue
                if concluse is False and g["stato"] in STATI_CONCLUSI:
                    continue
                out.append(riassunto(g))
    out.sort(key=lambda r: r["aggiornata"], reverse=True)
    return out


def scadenze(giorni: int | None = 7, mese: str | None = None) -> dict:
    """Gare non concluse con scadenza entro `giorni` (o nel mese 'YYYY-MM'), divise per settore."""
    oggi = datetime.now().date()
    tutte = elenco(concluse=False)
    sel = []
    for r in tutte:
        if not r["scadenza"]:
            continue
        try:
            d = datetime.fromisoformat(r["scadenza"]).date()
        except ValueError:
            continue
        if mese:
            if d.strftime("%Y-%m") == mese:
                sel.append(r)
        else:
            if oggi <= d <= oggi + timedelta(days=giorni or 7) or d < oggi:
                sel.append(r)
    sel.sort(key=lambda r: r["scadenza"])
    return {
        "cultura": [r for r in sel if r["settore"] == "Cultura"],
        "sociale": [r for r in sel if r["settore"] == "Sociale"],
        "altro": [r for r in sel if r["settore"] == "Altro"],
        "oggi": oggi.isoformat(),
    }


# ---------------------------------------------------------------------------
# Documenti
# ---------------------------------------------------------------------------

def estrai_testo(nome: str, dati: bytes) -> str:
    ext = Path(nome).suffix.lower()
    try:
        if ext == ".pdf":
            from io import BytesIO
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(dati))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if ext == ".docx":
            from io import BytesIO
            import docx
            d = docx.Document(BytesIO(dati))
            parti = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parti.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parti)
        if ext in (".txt", ".md", ".csv"):
            return dati.decode("utf-8", errors="replace")
    except Exception as e:  # noqa: BLE001
        return f"[Impossibile estrarre il testo da {nome}: {e}]"
    return f"[Formato {ext or 'sconosciuto'} non supportato per l'estrazione del testo]"


def aggiungi_documento(gid: str, nome: str, dati: bytes, categoria: str = "altro") -> dict:
    g = _carica_raw(gid)
    d = _dir(gid)
    (d / "documenti").mkdir(parents=True, exist_ok=True)
    (d / "testi").mkdir(parents=True, exist_ok=True)
    doc_id = uuid.uuid4().hex[:8]
    sicuro = re.sub(r"[^\w.\-]+", "_", nome)[:120]
    (d / "documenti" / f"{doc_id}_{sicuro}").write_bytes(dati)
    testo = estrai_testo(nome, dati)
    (d / "testi" / f"{doc_id}.txt").write_text(testo[:MAX_CHARS_DOC], encoding="utf-8")
    g["documenti"].append({
        "id": doc_id, "nome": nome, "categoria": categoria if categoria in CATEGORIE_DOC else "altro",
        "caricato": _now(), "byte": len(dati), "caratteri": min(len(testo), MAX_CHARS_DOC),
        "troncato": len(testo) > MAX_CHARS_DOC,
    })
    return tocca(g)


def rimuovi_documento(gid: str, doc_id: str) -> dict:
    g = _carica_raw(gid)
    d = _dir(gid)
    for f in (d / "documenti").glob(f"{doc_id}_*") if (d / "documenti").exists() else []:
        f.unlink()
    t = d / "testi" / f"{doc_id}.txt"
    if t.exists():
        t.unlink()
    g["documenti"] = [x for x in g["documenti"] if x["id"] != doc_id]
    return tocca(g)


def testo_documento(gid: str, doc_id: str) -> str:
    t = _dir(gid) / "testi" / f"{doc_id}.txt"
    return t.read_text(encoding="utf-8") if t.exists() else ""


def tutti_i_documenti() -> list[dict]:
    out = []
    for r in elenco():
        g = _carica_raw(r["id"])
        for doc in g["documenti"]:
            out.append({**doc, "gara_id": g["id"], "gara_titolo": g["titolo"], "settore": g["settore"]})
    out.sort(key=lambda x: x["caricato"], reverse=True)
    return out


# ---------------------------------------------------------------------------
# Contesto per l'assistente
# ---------------------------------------------------------------------------

def contesto_gara(g: dict) -> str:
    """Testo che descrive la gara, i suoi documenti e la memoria (se attiva)."""
    parti = [
        "# Gara in lavorazione",
        f"- Titolo: {g['titolo']}",
        f"- Ente: {g['ente'] or 'n.d.'}",
        f"- Settore: {g['settore']}",
        f"- Scadenza: {g.get('scadenza') or 'n.d.'}",
        f"- Stato: {g['stato']}",
        f"- Base d'asta: {g.get('base_asta') or 'n.d.'}",
    ]
    if g.get("note"):
        parti.append(f"- Note: {g['note']}")

    if g.get("memoria_attiva", True):
        if g.get("info_estratte"):
            parti.append("\n## Informazioni già estratte\n```json\n" + json.dumps(g["info_estratte"], ensure_ascii=False, indent=2) + "\n```")
        if g.get("valutazione"):
            parti.append("\n## Valutazione già effettuata\n" + g["valutazione"])
        if g.get("dati_simulatore"):
            parti.append("\n## Dati già estratti per il simulatore\n```json\n" + json.dumps(g["dati_simulatore"], ensure_ascii=False, indent=2) + "\n```")
    else:
        parti.append("\n(La memoria di questa gara è archiviata: analisi e conversazioni precedenti non sono nel contesto.)")

    budget = MAX_CHARS_CONTESTO
    parti.append("\n# Documenti della gara")
    if not g["documenti"]:
        parti.append("(nessun documento caricato)")
    for doc in g["documenti"]:
        testo = testo_documento(g["id"], doc["id"])
        if budget <= 0:
            parti.append(f"\n## {doc['nome']} [{doc['categoria']}]\n(omesso: limite di contesto raggiunto)")
            continue
        pezzo = testo[:budget]
        budget -= len(pezzo)
        parti.append(f"\n## {doc['nome']} [{doc['categoria']}]\n{pezzo}")
    return "\n".join(parti)


def messaggi_per_modello(g: dict, ultimi: int = 30) -> list[dict]:
    if not g.get("memoria_attiva", True):
        return []
    return [{"role": m["role"], "content": m["content"]} for m in g["chat"][-ultimi:]]


def aggiungi_messaggio(g: dict, role: str, content: str) -> dict:
    g["chat"].append({"role": role, "content": content, "ts": _now()})
    return tocca(g)


# ---------------------------------------------------------------------------
# Estrazione euristica dai PDF (fallback senza AI)
# ---------------------------------------------------------------------------

_RE_DATA = re.compile(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})")
_RE_ORA = re.compile(r"ore\s*(\d{1,2})[:.](\d{2})", re.IGNORECASE)
_KW_CULTURA = ("cultur", "museo", "musei", "bibliotec", "teatr", "archiv", "mostra", "patrimonio", "turis")
_KW_SOCIALE = ("social", "assistenz", "educativ", "minori", "anzian", "disabil", "welfare", "domicil", "asilo", "comunit")


def estrai_scadenza_euristica(testo: str, nome_file: str = "") -> dict:
    righe = [r.strip() for r in testo.splitlines() if r.strip()]
    titolo = ""
    for r in righe[:40]:
        if 20 < len(r) < 200 and any(k in r.lower() for k in ("affidamento", "procedura", "gara", "servizio", "appalto", "concessione")):
            titolo = r
            break
    if not titolo:
        titolo = Path(nome_file).stem.replace("_", " ") if nome_file else (righe[0][:150] if righe else "")

    ente = ""
    for r in righe[:60]:
        if any(k in r.lower() for k in ("comune di", "regione", "provincia di", "città metropolitana", "azienda", "asl", "ministero", "fondazione", "unione")):
            ente = r[:150]
            break

    data, ora = None, None
    testo_l = testo.lower()
    idx = -1
    for kw in ("termine per la presentazione", "termine di presentazione", "scadenza", "entro le ore", "termine ultimo", "presentazione delle offerte"):
        idx = testo_l.find(kw)
        if idx >= 0:
            break
    finestra = testo[idx: idx + 500] if idx >= 0 else testo[:5000]
    m = _RE_DATA.search(finestra) or _RE_DATA.search(testo)
    if m:
        gg, mm, aa = m.groups()
        try:
            data = datetime(int(aa), int(mm), int(gg)).date().isoformat()
        except ValueError:
            data = None
    mo = _RE_ORA.search(finestra) or _RE_ORA.search(testo)
    if mo:
        ora = f"{int(mo.group(1)):02d}:{mo.group(2)}"

    settore = "Altro"
    sc = sum(testo_l.count(k) for k in _KW_CULTURA)
    ss = sum(testo_l.count(k) for k in _KW_SOCIALE)
    if sc or ss:
        settore = "Cultura" if sc >= ss else "Sociale"

    return {"titolo": titolo, "ente": ente, "data_scadenza": data, "ora_scadenza": ora, "settore": settore,
            "base_asta": None, "note": "", "fonte": "euristica"}
